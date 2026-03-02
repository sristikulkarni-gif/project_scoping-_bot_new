import os
import tempfile
import anyio
import pytesseract
import openpyxl
import logging
from io import BytesIO
from typing import List
from PIL import Image
from docx import Document
from pptx import Presentation
from pdfminer.high_level import extract_text as extract_pdf_text

from app.utils import azure_blob

logger = logging.getLogger(__name__)

def extract_text_from_file(file_bytes_io: BytesIO, file_name: str) -> str:
    """
    Extract text from a file given its bytes and filename.
    """
    suffix = os.path.splitext(file_name)[-1].lower()
    file_bytes = file_bytes_io.read()
    file_bytes_io.seek(0)

    content = ""
    try:
        if suffix == ".pdf":
            with tempfile.NamedTemporaryFile(delete=False, suffix=suffix) as tmp:
                tmp.write(file_bytes)
                tmp_path = tmp.name
            try:
                content = extract_pdf_text(tmp_path)
            finally:
                os.remove(tmp_path)

        elif suffix == ".docx":
            doc = Document(BytesIO(file_bytes))
            content = "\n".join(p.text for p in doc.paragraphs)

        elif suffix == ".pptx":
            prs = Presentation(BytesIO(file_bytes))
            texts = []
            for slide in prs.slides:
                for shape in slide.shapes:
                    if hasattr(shape, "text"):
                        texts.append(shape.text)
            content = "\n".join(texts)

        elif suffix in [".xlsx", ".xlsm"]:
            wb = openpyxl.load_workbook(BytesIO(file_bytes))
            sheet = wb.active
            content = "\n".join(
                " ".join(str(cell) if cell else "" for cell in row)
                for row in sheet.iter_rows(values_only=True)
            )

        elif suffix in [".png", ".jpg", ".jpeg", ".tiff"]:
            img = Image.open(BytesIO(file_bytes))
            content = pytesseract.image_to_string(img)

        else:
            with tempfile.NamedTemporaryFile(delete=False, suffix=suffix) as tmp:
                tmp.write(file_bytes)
                tmp_path = tmp.name
            try:
                with open(tmp_path, "r", encoding="utf-8", errors="ignore") as fh:
                    content = fh.read()
            finally:
                os.remove(tmp_path)

    except Exception as e:
        logger.warning(f"Text extraction failed for {file_name}: {e}")

    return content.strip()

async def extract_text_from_files(files: List[dict]) -> str:
    results: List[str] = []

    async def _extract_single(f: dict) -> None:
        try:
            blob_bytes = await azure_blob.download_bytes(f["file_path"])
            suffix = os.path.splitext(f["file_name"])[-1].lower()

            def process_file() -> str:
                content = ""
                try:
                    if suffix == ".pdf":
                        with tempfile.NamedTemporaryFile(delete=False, suffix=suffix) as tmp:
                            tmp.write(blob_bytes)
                            tmp_path = tmp.name
                        try:
                            content = extract_pdf_text(tmp_path)
                        finally:
                            os.remove(tmp_path)

                    elif suffix == ".docx":
                        doc = Document(BytesIO(blob_bytes))
                        content = "\n".join(p.text for p in doc.paragraphs)

                    elif suffix == ".pptx":
                        prs = Presentation(BytesIO(blob_bytes))
                        texts = []
                        for slide in prs.slides:
                            for shape in slide.shapes:
                                if hasattr(shape, "text"):
                                    texts.append(shape.text)
                        content = "\n".join(texts)

                    elif suffix in [".xlsx", ".xlsm"]:
                        wb = openpyxl.load_workbook(BytesIO(blob_bytes))
                        sheet = wb.active
                        content = "\n".join(
                            " ".join(str(cell) if cell else "" for cell in row)
                            for row in sheet.iter_rows(values_only=True)
                        )

                    elif suffix in [".png", ".jpg", ".jpeg", ".tiff"]:
                        img = Image.open(BytesIO(blob_bytes))
                        content = pytesseract.image_to_string(img)

                    else:
                        with tempfile.NamedTemporaryFile(delete=False, suffix=suffix) as tmp:
                            tmp.write(blob_bytes)
                            tmp_path = tmp.name
                        try:
                            with open(tmp_path, "r", encoding="utf-8", errors="ignore") as fh:
                                content = fh.read()
                        finally:
                            os.remove(tmp_path)

                except Exception as e:
                    logger.warning(f"Extraction failed for {f['file_name']}: {e}")

                return content.strip()

            text = await anyio.to_thread.run_sync(process_file)

            if text:
                results.append(text)
            else:
                logger.warning(f"Extracted no text from {f['file_name']}")

        except Exception as e:
            logger.warning(f"Failed to extract {f.get('file_name')} (path={f.get('file_path')}): {e}")

    async with anyio.create_task_group() as tg:
        for f in files:
            tg.start_soon(_extract_single, f)

    return "\n\n".join(results)

def extract_document_overview(full_text: str, max_chars: int = 10000) -> str:
    if not full_text:
        return ""

    if len(full_text) <= max_chars:
        return full_text

    overview = full_text[:max_chars]

    last_period = overview.rfind('. ')
    if last_period > max_chars * 0.8:
        overview = overview[:last_period + 1]

    return overview
